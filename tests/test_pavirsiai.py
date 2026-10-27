"""Paviršių palyginimas: sintetiniai duomenys su žinomu atsakymu (TTM, poslinkis, duobė, ruožai)."""
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.geo import io_las, io_ttm, loaders  # noqa: E402
from app.geo.polyline import centerline, station_offset  # noqa: E402
from app.geo.surface import Grid, find_zones  # noqa: E402

E0, N0 = 522000.0, 6127000.0


def write_ttm(path, pts, tris):
    """Minimalus TTM (kaip TBC): 4 B koordinatės ir aukščiai, 4 B viršūnių nr."""
    on, oe = pts[:, 1].min(), pts[:, 0].min()
    h = bytearray(b"\x00" * 160)
    h[0:2] = b"\x01\x00"
    h[2:18] = io_ttm.SIGNATURE
    h[22:26] = b"TEST"
    h[54] = 1
    struct.pack_into("<6d", h, 57, on, oe, 0, 0, 0, 0)
    h[105:109] = bytes([4, 4, 4, 4])
    nv, nt = len(pts), len(tris)
    sv = 160
    st = sv + nv * 12
    struct.pack_into("<iihiih", h, 109, sv, nv, 12, st, nt, 24)
    v = b"".join(struct.pack("<3f", p[1] - on, p[0] - oe, p[2]) for p in pts)
    t = b"".join(struct.pack("<6i", a + 1, b + 1, c + 1, 0, 0, 0) for a, b, c in tris)
    Path(path).write_bytes(bytes(h) + v + t)


def design_strip():
    """Juosta 6 x 120 m (išilgai N), nuolydis skersai 2 %."""
    xs, ys = np.arange(0, 6.01, 1.0), np.arange(0, 120.01, 5.0)
    X, Y = np.meshgrid(xs, ys)
    pts = np.column_stack([X.ravel() + E0, Y.ravel() + N0, 100 + 0.02 * X.ravel()])
    nx = len(xs)
    tris = []
    for j in range(len(ys) - 1):
        for i in range(nx - 1):
            a = j * nx + i
            tris += [(a, a + 1, a + nx + 1), (a, a + nx + 1, a + nx)]
    return pts, np.array(tris)


def test_ttm_roundtrip(tmp):
    pts, tris = design_strip()
    f = tmp / "d.ttm"
    write_ttm(f, pts, tris)
    d = io_ttm.read(f)
    assert len(d["pts"]) == len(pts) and len(d["tris"]) == len(tris)
    assert np.abs(d["pts"] - pts).max() < 0.002
    S = loaders.load_surface(f)
    assert abs(S.z_at([[E0 + 3, N0 + 50]])[0] - 100.06) < 1e-3
    assert abs(S.area_2d() - 720) < 0.01


def test_zones():
    z = np.zeros((40, 40))
    z[5:9, 5:13] = -0.10          # 1 x 2 m (žingsnis 0.25)
    z[20:22, 20:22] = -0.06       # 0.25 m² - mažesnė už min plotą
    z[30, 30] = np.nan
    zs = find_zones(Grid(0.125, 0.125, 0.25, z), 0.05, 0.5)
    assert len(zs) == 1
    assert abs(zs[0].area - 2.0) < 1e-9 and abs(zs[0].volume - 0.2) < 1e-9 and abs(zs[0].max_dz - 0.1) < 1e-9


def test_centerline():
    t = np.linspace(0, np.pi / 2, 2000)
    R = 200
    r = R + np.random.default_rng(1).uniform(-3, 3, len(t))
    P = np.column_stack([r * np.cos(t), r * np.sin(t)])
    C = centerline(P, seg=5)
    st, of = station_offset(C, P)
    assert np.ptp(st) > 0.95 * R * np.pi / 2
    assert np.abs(of).max() < 3.6


def test_tool_cli(tmp):
    pts, tris = design_strip()
    write_ttm(tmp / "d.ttm", pts, tris)
    rng = np.random.default_rng(0)
    n = 300_000
    xy = np.column_stack([rng.uniform(0, 6, n), rng.uniform(0, 120, n)])
    z = 100 + 0.02 * xy[:, 0] + 0.030 + rng.normal(0, 0.005, n)           # GPS +30 mm, triukšmas 5 mm
    pit = (xy[:, 0] > 2) & (xy[:, 0] < 4) & (xy[:, 1] > 40) & (xy[:, 1] < 45)
    z[pit] -= 0.15                                                     # išplova 2 x 5 m, 150 mm = 1.5 m³
    io_las.write(tmp / "s.las", np.column_stack([xy[:, 0] + E0, xy[:, 1] + N0, z]))
    out = tmp / "out"
    r = subprocess.run([sys.executable, str(ROOT / "irankine.py"), "run", "pavirsiai", f"design={tmp / 'd.ttm'}",
                        f"scan={tmp / 's.las'}", "--set", "fit_slope=0", "--set", "zone_thr=50,100",
                        "--out", str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    import openpyxl
    wb = openpyxl.load_workbook(out / "d_ataskaita.xlsx")
    info = {row[0].value: row[1].value for row in wb["Zonos"].iter_rows(max_col=2) if row[0].value}
    cut = info["Žemiau projekto (iškasa / išplovos), m³"]
    fill = info["Aukščiau projekto (pylimas), m³"]
    assert abs(cut - 1.5) < 0.08, cut
    assert fill < 0.05, fill
    assert "+30" in info["Pritaikymas"] or "+29" in info["Pritaikymas"] or "+31" in info["Pritaikymas"]
    assert "Ruožai" in wb.sheetnames                                  # ilgas objektas
    for f in ("d_palyginimas.dxf", "d_dz.png", "d_dz.pgw", "d_pritaikytas.xml", "d_planas.png"):
        assert (out / f).exists(), f
    assert "A1" in r.stdout and "B1" in r.stdout


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        for name, fn in list(globals().items()):
            if name.startswith("test_"):
                fn(tmp) if fn.__code__.co_argcount else fn()
                print("OK", name)
