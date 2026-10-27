"""Išpildomųjų taškų kontrolė: žinomi dz, kodų nuokrypiai, PK / atstumas, keli paviršiai."""
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.geo import io_landxml  # noqa: E402
from app.geo.alignment import Alignment  # noqa: E402

E0, N0 = 500000.0, 6100000.0


def make(tmp):
    # ašis į šiaurę nuo (E0+5, N0), 100 m; danga 10 m pločio, skersinis nuolydis 2.5 % į dešinę (rytus) žemyn
    xs, ys = np.arange(0, 10.01, 1.0), np.arange(0, 100.01, 5.0)
    X, Y = np.meshgrid(xs, ys)
    nx = len(xs)
    tris = []
    for j in range(len(ys) - 1):
        for i in range(nx - 1):
            a = j * nx + i
            tris += [(a, a + 1, a + nx + 1), (a, a + nx + 1, a + nx)]

    def surf(name, dzz):
        z = 50 + 0.01 * Y.ravel() - 0.025 * (X.ravel() - 5) + dzz
        return {"name": name, "pts": np.column_stack([X.ravel() + E0, Y.ravel() + N0, z]), "tris": np.array(tris)}
    io_landxml.write_surfaces(tmp / "proj.xml", [surf("Danga", 0.0), surf("Pagrindas", -0.10)])
    al = Alignment.from_elems("ASIS", [("line", 100, 0, 0)], (E0 + 5, N0), np.pi / 2)
    io_landxml.write_alignments(tmp / "asis.xml", [al])
    # taškai: AS ant dangos +dz, PAG ant pagrindo
    pts = []
    rng = np.random.default_rng(3)
    for k in range(40):
        e, n = rng.uniform(0.5, 9.5), rng.uniform(1, 99)
        dz = [0.004, -0.008, 0.012, -0.015][k % 4]          # AS: ±10 -> 2 iš 4 blogi (12, -15)
        pts.append((f"A{k}", n + N0, e + E0, 50 + 0.01 * n - 0.025 * (e - 5) + dz, "AS"))
    for k in range(20):
        e, n = rng.uniform(0.5, 9.5), rng.uniform(1, 99)
        dz = [0.015, -0.025][k % 2]                           # PAG: 20/30 -> visi OK
        pts.append((f"P{k}", n + N0, e + E0, 50 + 0.01 * n - 0.025 * (e - 5) - 0.10 + dz, "PAG"))
    pts.append(("X1", N0 + 50, E0 + 30, 50.0, "AS"))            # už ribų
    pts.append(("B1", N0 + 50, E0 + 5, 99.0, "BAZE"))           # praleidžiamas
    with open(tmp / "tsk.csv", "w") as fh:
        for r in pts:
            fh.write(f"{r[0]},{r[1]:.3f},{r[2]:.3f},{r[3]:.4f},{r[4]}\n")


def test_cli(tmp):
    make(tmp)
    out = tmp / "out"
    r = subprocess.run([sys.executable, str(ROOT / "irankine.py"), "run", "kontrole", str(tmp / "tsk.csv"),
                        str(tmp / "proj.xml"), str(tmp / "asis.xml"),
                        "--set", "tol=10", "--set", "tol_codes=PAG:20/30", "--set", "surf_map=PAG=Pagrindas",
                        "--set", "skip_codes=BAZE", "--out", str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    import openpyxl
    ws = openpyxl.load_workbook(out / "tsk_kontrole.xlsx")["Suvestinė"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    hdr = next(i for i, row in enumerate(rows) if row[0] == "Kodas")
    H = rows[hdr]
    tab = {row[0]: dict(zip(H, row)) for row in rows[hdr + 1:] if row[0]}
    assert tab["AS"]["Taškų"] == 40 and tab["AS"]["Atitiktis, %"] == 50.0, tab["AS"]
    assert tab["AS"]["Per aukštai"] == 10 and tab["AS"]["Per žemai"] == 10
    assert tab["PAG"]["Atitiktis, %"] == 100.0 and abs(tab["PAG"]["Vidurkis, mm"] + 5) < 0.6, tab["PAG"]
    assert "BAZE" not in tab
    pts = openpyxl.load_workbook(out / "tsk_kontrole.xlsx")["Taškai"]
    rows = [[c.value for c in row] for row in pts.iter_rows()]
    hdr = next(i for i, row in enumerate(rows) if row[0] == "Nr")
    H = rows[hdr]
    P = {row[0]: dict(zip(H, row)) for row in rows[hdr + 1:] if row[0]}
    x1 = P["X1"]
    assert x1["Būsena"] == "už ribų"
    assert abs(x1["PK"] - 50) < 1e-3 and abs(x1["Atstumas nuo ašies, m (+ dešinėje)"] - 25) < 1e-3   # rytai = dešinė
    assert "Už paviršiaus ribų: 1" in r.stdout and "NEATITINKA" in r.stdout
    for f in ("tsk_kontrole.dxf", "tsk_dz.csv", "tsk_kontrole.png"):
        assert (out / f).exists()


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as d:
        test_cli(Path(d))
        print("OK test_cli")
