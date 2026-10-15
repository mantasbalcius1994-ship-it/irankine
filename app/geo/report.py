"""Excel ataskaitos (openpyxl): antraštės informacija, lentelės su formatais, žymėjimas už tolerancijos."""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1D3550")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BAD_FILL = PatternFill("solid", fgColor="FCE4EC")
BAD_FONT = Font(color="B00040", bold=True)
THIN = Side(style="thin", color="D6DDE4")


@dataclass
class Col:
    title: str
    key: str | Callable[[dict], Any]
    fmt: str | None = None          # pvz. "0.000", "0.0%", "0"
    width: float | None = None


@dataclass
class Sheet:
    name: str
    cols: list[Col]
    rows: list[dict]
    bad: Callable[[dict], bool] | None = None     # eilutė paryškinama
    bad_cols: list[str] = field(default_factory=list)   # kurie stulpeliai raudoni (title)
    info: list[tuple[str, Any]] = field(default_factory=list)  # [(pavadinimas, reikšmė)] virš lentelės
    totals: dict[str, Any] = field(default_factory=dict)       # {stulpelio title: reikšmė/formulė 'SUM'}


def write_xlsx(path, sheets: list[Sheet], title: str | None = None):
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh.name[:31])
        r = 1
        if title:
            ws.cell(r, 1, title).font = Font(bold=True, size=13); r += 1
        for k, v in [("Sukurta", dt.datetime.now().strftime("%Y-%m-%d %H:%M"))] + list(sh.info):
            ws.cell(r, 1, k).font = Font(color="5B6B7A")
            ws.cell(r, 2, v)
            r += 1
        r += 1
        head = r
        for j, c in enumerate(sh.cols, 1):
            cell = ws.cell(r, j, c.title)
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30
        r += 1
        first = r
        for row in sh.rows:
            bad = bool(sh.bad and sh.bad(row))
            for j, c in enumerate(sh.cols, 1):
                v = c.key(row) if callable(c.key) else row.get(c.key)
                cell = ws.cell(r, j, v)
                if c.fmt:
                    cell.number_format = c.fmt
                cell.border = Border(bottom=THIN)
                if bad:
                    cell.fill = BAD_FILL
                    if c.title in sh.bad_cols:
                        cell.font = BAD_FONT
            r += 1
        if sh.totals and sh.rows:
            for j, c in enumerate(sh.cols, 1):
                if c.title in sh.totals:
                    v = sh.totals[c.title]
                    L = get_column_letter(j)
                    if v in ("SUM", "MAX", "MIN", "AVERAGE"):
                        v = f"={v}({L}{first}:{L}{r - 1})"
                    cell = ws.cell(r, j, v)
                    cell.font = Font(bold=True)
                    if c.fmt:
                        cell.number_format = c.fmt
        for j, c in enumerate(sh.cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = c.width or max(10, min(40, len(c.title) + 4))
        ws.freeze_panes = ws.cell(head + 1, 1)
        if sh.rows:
            ws.auto_filter.ref = f"A{head}:{get_column_letter(len(sh.cols))}{r - 1}"
    wb.save(path)
